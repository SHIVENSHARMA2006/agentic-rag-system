from pathlib import Path

import fitz
import pandas as pd
from docx import Document
from pptx import Presentation


class ParserService:

    def parse(self, file_path: str) -> str:

        suffix = Path(file_path).suffix.lower()

        if suffix == ".pdf":
            return self._parse_pdf(file_path)

        if suffix == ".docx":
            return self._parse_docx(file_path)

        if suffix == ".pptx":
            return self._parse_pptx(file_path)

        if suffix == ".xlsx":
            return self._parse_excel(file_path)

        if suffix in [".txt", ".md", ".csv"]:
            return Path(file_path).read_text(
                encoding="utf-8",
                errors="ignore",
            )

        raise ValueError(f"Unsupported file type: {suffix}")

    def _parse_pdf(self, file_path):

        doc = fitz.open(file_path)

        text = ""

        for page in doc:
            text += page.get_text()

        return text

    def _parse_docx(self, file_path):

        doc = Document(file_path)

        return "\n".join(
            p.text for p in doc.paragraphs
        )

    def _parse_pptx(self, file_path):

        prs = Presentation(file_path)

        text = ""

        for slide in prs.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text += shape.text + "\n"

        return text

    def _parse_excel(self, file_path):

        sheets = pd.read_excel(
            file_path,
            sheet_name=None,
        )

        text = ""

        for _, df in sheets.items():

            text += df.to_string()

        return text